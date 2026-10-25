from pathlib import Path
import hashlib
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).parent
SOURCE = ROOT / "source" / "REG-ARCH_v0_4.docx"
OUT = ROOT / "REG-ARCH_v0.4.1-candidate.docx"
SOURCE_PDF = ROOT / "source" / "REG-ARCH_v0_4.pdf"
SOURCE_HASH = "46752A0E1CD7AAE430742878C14C6B5C28FE1DA8167C9A81D78B23D9BAA7D3C2"
SOURCE_PDF_HASH = "4BE9D68CA8DCE3557655905467C0975758E82AB2888F974B0B9439D59620B939"

def replace_paragraph(p, text):
    p.clear()
    p.add_run(text)

if hashlib.sha256(SOURCE.read_bytes()).hexdigest().upper() != SOURCE_HASH:
    raise SystemExit("Source DOCX hash mismatch; candidate not built")
if hashlib.sha256(SOURCE_PDF.read_bytes()).hexdigest().upper() != SOURCE_PDF_HASH:
    raise SystemExit("Source PDF hash mismatch; candidate not built")
doc = Document(SOURCE)
paras = doc.paragraphs

# The source version is in the metadata table, not in a paragraph.
for row in doc.tables[0].rows:
    if row.cells[0].text.strip() == "Version":
        row.cells[1].text = "0.4.1 candidate (non ratifiée)"
    if row.cells[0].text.strip() == "Intention":
        row.cells[1].text = "Amendement de cohérence normative ; proposition non ratifiée"
    if row.cells[0].text.strip() == "Statut":
        row.cells[1].text = "Candidate normative — non ratifiée, non promue"
    if row.cells[0].text.strip() == "Source":
        row.cells[1].text = "REG-ARCH v0.4 + amendements P0 de la candidate v0.4.1"
    if row.cells[0].text.strip() == "Périmètre":
        row.cells[1].text = "Modèle architectural, raffinement borné et traçabilité PREPARE/PREVENT"

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if "Machines versionnées + porte G-Écart" in cell.text:
                cell.text = cell.text.replace("Machines versionnées + porte G-Écart", "Machines versionnées des cycles LC ; G-Écart défini par REG-ARCH-025")
            if "Quels états et transitions sont canoniques pour chaque cycle ?" in cell.text:
                cell.text = "Quels états et transitions complets sont canoniques pour chaque cycle LC ? G-Écart est déjà défini par REG-ARCH-025."
            if "Inclut explicitement la porte G-Écart parmi les artefacts attendus" in cell.text:
                cell.text = "Pour v0.4 : G-Écart encore attendue ; cette lacune est corrigée par REG-ARCH-025 en v0.4.1 candidate. D-02 reste ouverte pour les automates LC."
            if "Note sur la porte de création d'Écart : non encore définie" in cell.text:
                cell.text = "État initial v0.4 (non définie, transition informelle) — contradiction supprimée dans la candidate par REG-ARCH-025."

# Update document identity and history while preserving the source document's tables/layout.
for p in paras:
    if p.text.startswith("REG-ARCH v0.4 — Spécification d’architecture") or p.text.startswith("REG-ARCH v0.4 — Spécification d'architecture"):
        replace_paragraph(p, "REG-ARCH v0.4.1 candidate — Spécification d’architecture amendée")
    elif p.text.startswith("Version 0.4 — Document autonome"):
        replace_paragraph(p, "Version 0.4.1 candidate — proposition amendée non ratifiée")
    elif p.text.startswith("Ajouts propres à la v0.4 :"):
        replace_paragraph(p, p.text + "\nAjouts propres à la candidate v0.4.1 : porte G-Écart testable ; contrat de raffinement borné ; obligations PREPARE/PREVENT, invalidation transitive et requalification ; appendice explicitement informatif.")
    elif p.text.startswith("REG-ARCH-010bis —"):
        replace_paragraph(p, "REG-ARCH-010bis — Refus structurel et création d’un Écart distinct. Un refus est structurel uniquement s’il résulte (a) de l’absence ou de l’invalidité d’une autorité requise, (b) d’une liaison insuffisante ou ambiguë entre objets, ou (c) de critères normatifs absents ou indéterminés sur l’objet source. Dans ces seuls cas, le traitement de la cible DOIT être refusé et sa version, son état et ses événements antérieurs DOIVENT rester inchangés. Le système DOIT ensuite tenter, au moyen de G-Écart définie en REG-ARCH-025 (§17), la création d’un nouvel objet Écart LC-08 ; cet objet n’est pas l’objet cible et sa création ne transforme pas le refus en acceptation. Un refus opérationnel ordinaire (dont clé consommée, condition d’autorisation non satisfaite alors que l’autorité et les critères sont valides, ou issue externe indéterminée) ne déclenche pas cette règle. Si la trace durable de refus ou le stockage requis manque, aucune application cible ni création d’Écart n’est permise ; le résultat ne peut être qualifié de traité. La création G-Écart doit appliquer le contrat d’atomicité et de reprise défini en REG-ARCH-025 ; une transition informelle est interdite.")
    elif p.text.startswith("Note sur la porte de création d'Écart") or p.text.startswith("Note sur la porte de création d’Écart"):
        replace_paragraph(p, "La transition G-Écart est définie normativement par REG-ARCH-025 (§17) et ne dépend pas de la clôture de D-02. D-02 reste ouverte pour les automates complets des cycles LC.")
    elif p.text.startswith("Fin de la spécification REG-ARCH v0.4."):
        replace_paragraph(p, "Fin du corps normatif REG-ARCH v0.4.1 candidate. Cette candidate n’est ni ratifiée ni promue.")

for section in doc.sections:
    for p in section.header.paragraphs + section.footer.paragraphs:
        if "REG-ARCH v0.4 — Spécification d’architecture" in p.text or "REG-ARCH v0.4 — Spécification d'architecture" in p.text:
            replace_paragraph(p, "REG-ARCH v0.4.1 candidate — Spécification d’architecture")

# Insert normative clauses before the original appendix, after the existing change history.
marker = next(p for p in doc.paragraphs if p.text.startswith("Fin du corps normatif REG-ARCH v0.4.1 candidate."))
new = [
    ("heading", "17. Amendements normatifs v0.4.1 candidate"),
    ("body", "Les obligations REG-ARCH-025 à 027 complètent les obligations 001 à 024 et 010bis. Elles sont normatives pour cette candidate. Elles n’imposent ni langage, ni base de données, ni moteur, ni topologie d’implémentation. Les termes DOIT, NE DOIT PAS et PEUT ont leur sens normatif habituel dans ce document."),
    ("heading", "REG-ARCH-025 — Porte G-Écart"),
    ("body", "Autorité et identité : G-Écart est une transition réservée de tenue de registre, définie par REG-ARCH et déclenchée uniquement par le prédicat de refus structurel de 010bis. Son acteur est le contrôleur de transition du registre, sans délégation ni identité humaine à approuver ; son mandat est limité à créer ou retrouver l’Écart lié au refus et ne permet ni admission, ni modification de la cible. Le contrôleur DOIT utiliser la version de REG-ARCH déclarée active par le manifeste normatif du périmètre. L’identité effective du contrôleur DOIT être liée à une assertion de capacité versionnée, vérifiable par le point d’enforcement du registre : capability_id=record_structural_refusal, portée=création/résolution d’Écart seulement, aucune permission d’admission ou mutation cible, référence de principal et de liaison de confiance inscrites dans la trace. Si principal, assertion, portée, liaison de confiance, version active ou schéma d’événement ne peut être vérifié, G-Écart refuse et ne peut déclarer le refus traité. Cette autorité ne résout pas D-01, qui porte sur l’autorité métier d’admission."),
    ("body", "Schéma de trace obligatoire : chaque invocation traitée par le journal durable émet un événement versionné avec les champs obligatoires `event_id`, `event_type`, `schema_version`, `request_id`, `source_refusal_event_id`, `causation_event_ids`, `source_object_id`, `source_object_version`, `target_object_id`, `target_object_version`, `refusal_category`, `reason_code`, `gate_id`, `gate_version`, `reg_arch_version`, `active_manifest_id`, `active_manifest_version`, `active_manifest_digest`, `requester_ref`, `controller_principal_ref`, `capability_assertion_ref`, `trust_binding_ref`, `capability_scope_version`, `logical_order`, `idempotency_key`, `outcome` et `gap_object_id`. `outcome` est exactement l’un de CREATED, REUSED, REFUSED, ABORTED ; `gap_object_id` est nul seulement pour REFUSED ou ABORTED. Une tentative refusée faute de contexte, de manifeste ou de liaison d’identité DOIT produire un événement REFUSED ; tout champ indisponible dans ce refus est présent avec valeur nulle et `reason_code` qui identifie la précondition non satisfaite. La nullité est interdite pour CREATED/REUSED. Si le journal durable lui-même est indisponible, le contrôleur NE DOIT PAS entrer dans G-Écart ni créer un objet ; le refus source demeure non traité et aucune réussite ne peut être annoncée. Aucun champ obligatoire n’est omis. Le schéma JSON associé fixe les types, nullabilités et énumérations."),
    ("body", "Préconditions et résultat : (1) un événement durable de refus structurel conforme à 010bis existe ; (2) ses versions, cible, source, motif et porte sont résolubles ; (3) la clé d’idempotence est déterministement dérivée de l’identifiant d’événement de refus et de la version de G-Écart ; (4) l’identité cible de l’Écart est distincte de l’objet cible. G-Écart refuse les refus ordinaires, causalité absente, versions non résolues, manifeste actif absent, clé collisionnelle avec contenu divergent ou schéma invalide. Si la trace durable ne peut être écrite, aucune application cible ni affirmation de traitement réussi n’est permise."),
    ("body", "Causalité, atomicité et reprise : chaque tentative référence l’événement source dans causation_event_ids. L’objet Écart et son événement de création sont visibles atomiquement : aucun objet sans événement complet ni événement de création sans objet. Un crash avant commit ne laisse aucun résultat partiel. Un crash après commit mais avant accusé de réception est résolu par rejeu de la même clé déterministe : le contrôleur retourne le même Écart déjà créé, sans doublon ni nouvel événement de création. En cas de collision, `request_fingerprint` DOIT être le SHA-256 (hexadécimal minuscule) des octets UTF-8 de la représentation JSON canonique JCS conforme à RFC 8785 de l’objet contenant exactement ces propriétés et valeurs : `source_refusal_event_id`, `source_object_id`, `source_object_version`, `target_object_id`, `target_object_version`, `refusal_category`, `requested_reason`, `reg_arch_version`, `gate_version`, `active_manifest_id`, `active_manifest_version`, `active_manifest_digest`, `requester_ref`, `controller_principal_ref`, `capability_assertion_ref`, `trust_binding_ref`, `capability_scope_version`. Une valeur non résolue est JSON null. Aucun champ supplémentaire n’est inclus. Si la même clé est présentée avec un contenu divergent, le refus DOIT produire un événement distinct construit à partir du contexte de la requête rejetée, citer l’événement original dans `conflicting_prior_event_id`, conserver le motif et l’empreinte ci-dessus, et utiliser `reason_code=IDEMPOTENCY_KEY_CONFLICT`; il ne doit pas copier les champs de contenu de l’événement initial comme s’ils décrivaient la requête rejetée. Un incident est lié à cette trace. L’état et la version de l’objet cible ne changent jamais. L’atomicité est une propriété observable exigée, sans prescrire la technologie de transaction."),
    ("body", "Oracle de conformité : tester refus structurel avec identité/versions et trace complètes, refus ordinaire, manifeste actif absent, causalité absente, collision divergente, crash avant commit, crash après commit avant ACK puis reprise avec la même clé. Le résultat accepté doit laisser la cible inchangée et produire exactement un Écart causalement résoluble ; la reprise doit retourner le même identifiant. Tout objet/événement partiel, doublon, application de cible ou traitement réussi sans trace est un échec."),
    ("heading", "REG-ARCH-026 — Raffinement et statut de l’appendice"),
    ("body", "L’appendice « How to validate Reg-Arch » est informatif : il n’ajoute ni obligation, ni garantie, ni équivalence. Les obligations normatives sont celles du corps identifié comme tel. Toute partie revendiquant qu’un système, profil ou configuration réalise une garantie REG-ARCH DOIT publier une déclaration de raffinement versionnée qui identifie (a) la garantie et les obligations couvertes, (b) le périmètre et les exclusions, (c) les versions exactes de REG-ARCH, REG, REG+, TENIR-Gov ou autres composants effectivement concernés, (d) les hypothèses et leur méthode de validation, (e) la relation de correspondance entre états/événements abstraits et traces concrètes, (f) les preuves et leur validité. Une garantie sans déclaration ou preuve suffisante est UNVERIFIABLE, et non réputée conforme."),
    ("body", "Critère de raffinement : pour chaque garantie revendiquée, la déclaration DOIT définir une relation d’abstraction versionnée entre états/événements concrets et modèle abstrait, un état initial correspondant, et les règles d’appariement des transitions. Chaque pas concret observable doit correspondre à un ou plusieurs pas abstraits autorisés ; un pas interne peut être ignoré seulement s’il ne modifie aucun état, événement ni propriété observable pertinente. Aucune transition concrète ne peut produire un effet interdit par l’obligation revendiquée. La preuve DOIT couvrir toutes les transitions de la garantie dans le périmètre, les refus, erreurs et conditions de terminaison pertinents, et lier chaque hypothèse à une vérification. Une preuve par échantillonnage seule ne démontre pas la couverture exhaustive d’un espace de transitions."),
    ("body", "Le raffinement est local aux garanties revendiquées. Aucune équivalence globale REG-ARCH → REG → REG+ → TENIR-Gov, aucune identité de composants et aucune implémentation particulière ne sont présumées. Un profil ou composant peut raffiner une partie, plusieurs parties, ou aucune ; l’absence de revendication ne vaut pas preuve. Une correspondance sans définition de relation, preuve de couverture des transitions et vérification des hypothèses est UNVERIFIABLE."),
    ("heading", "REG-ARCH-027 — PREPARE, PREVENT et invalidation"),
    ("body", "Complétude du graphe : chaque revendication DOIT inclure un manifeste fermé des dépendances nécessaires à la propriété, avec identifiants, versions/empreintes, relations typées et portée. La fermeture transitive doit résoudre chaque référence ; aucun nœud requis ne peut être absent, inconnu ou à version flottante. Le manifeste DOIT être comparé à un inventaire indépendant des entrées effectives de décision (règles, profil chargé, configuration, code/environnement et hypothèses observées) ; tout élément effectif non déclaré ou impossible à inventorier invalide la complétude et interdit SUPPORTED. La validité d’un arc exige sa justification et sa preuve, la simple présence dans le graphe n’est pas une preuve."),
    ("body", "PREPARE descendant : à partir d’une garantie et d’un instantané d’impact, résoudre la fermeture transitive, vérifier les identités et versions exactes, le manifeste fermé, la validité de chaque arc, les hypothèses, la portée et l’état de chaque preuve. Toute dépendance requise absente, incompatible, inconnue ou STALE interdit un verdict positif. Avant de publier un verdict, PREPARE DOIT vérifier que l’époque/version de l’instantané n’a pas changé ; si elle a changé, il abandonne l’évaluation et recommence sur le nouvel instantané."),
    ("body", "PREVENT ascendant : chaque changement observé crée une nouvelle version immuable et incrémente atomiquement une époque d’impact avant que la configuration modifiée devienne effective. L’incrément rend immédiatement toute preuve d’une époque antérieure non consommable, même si le calcul d’impact est différé. PREVENT marque ensuite atomiquement les preuves/verdicts transitifs dépendants STALE et produit les nœuds atteints, garanties affectées et dépendances manquantes. Un changement possède une clé d’opération stable ; une reprise après crash/ACK perdu retrouve la même époque et complète le calcul sans perte ni double changement. Si le calcul est incomplet, le périmètre concerné est fail closed."),
    ("body", "Protocole de consommation : aucune partie ne peut utiliser une preuve, même après PREPARE réussi, sans présenter un jeton lié à son époque, au digest du manifeste et aux versions dépendantes. Au point de consommation, le contrôle d’égalité avec l’époque courante, l’état non-STALE et l’enregistrement de la consommation DOIVENT partager un même point de linéarisation. Si le consommateur déclenche une décision gouvernée, celle-ci ne peut être publiée que dans la même époque validée. Toute modification ordonnée avant ce point fait refuser l’usage ; si la consommation est ordonnée avant la modification, l’historique doit conserver cet ordre. Une vérification séparée suivie d’un usage non synchronisé est non conforme."),
    ("body", "STALE et requalification : toute preuve et tout verdict dérivés d’une dépendance transitivement modifiée, révoquée ou non vérifiable deviennent STALE avant nouvel usage. STALE interdit soutien d’un verdict, promotion et livraison comme preuve courante. Une requalification exige une nouvelle preuve liée aux versions actuelles, nouvelle validation des hypothèses, manifeste complet et nouvel examen PREPARE ; elle ne réactive jamais l’ancien artefact."),
    ("body", "Oracle : (a) comparer un manifeste à l’inventaire effectif et injecter un élément non déclaré ; le résultat doit être incomplet/UNVERIFIABLE ; (b) modifier une configuration et vérifier la fermeture transitive attendue ; (c) garder l’identifiant lisible mais changer l’empreinte ; (d) exécuter PREPARE en concurrence avec la mutation et vérifier qu’aucun verdict ancien ne se publie ; (e) tenter la consommation après mutation entre vérification et usage et vérifier le refus au point de linéarisation ; (f) interrompre PREVENT après commit d’époque mais avant ACK, le reprendre avec la même clé et vérifier époque/impact stables ; (g) vérifier que la requalification crée une nouvelle preuve liée aux versions et manifeste actuels, tandis que l’ancienne reste STALE."),
    ("heading", "18. Décisions toujours ouvertes"),
    ("body", "D-01 à D-09 restent explicitement ouvertes dans cette candidate, sauf que la définition normative de G-Écart n’est plus différée à D-02. En particulier, cette candidate ne choisit pas l’autorité métier d’admission (D-01), les automates complets de tous cycles (D-02), les règles du guichet (D-03), les seuils et méthodes de mesure (D-04), la validation des approbations humaines (D-05), les garanties opérationnelles de persistance/concurrence/récupération (D-06), la rétention et les droits (D-07), la confiance cryptographique/horloges/ancrages (D-08), ni une validation réelle indépendante (D-09). Ces décisions limitent les garanties démontrables et doivent être déclarées comme hypothèses ou lacunes."),
]
for kind, text in new:
    p = marker.insert_paragraph_before(text)
    p.style = "Heading 2" if kind == "heading" else "Normal"

# Replace the v0.4 appendix wholesale: it contained product-specific assertions that
# cannot be made normative or reliable by a disclaimer alone.
body = doc._element.body
children = list(body.iterchildren())
marker_index = next(i for i, child in enumerate(children) if child is marker._p)
for child in children[marker_index + 1:]:
    if child.tag.endswith("}sectPr"):
        continue
    body.remove(child)
appendix = doc.add_paragraph("Appendix A — Informative, nonnormative")
appendix.style = "Heading 1"
doc.add_paragraph("Cet appendice n’ajoute aucune obligation, garantie, hiérarchie, équivalence ni prescription de produit. En cas de divergence, seuls le corps normatif et les obligations REG-ARCH-001 à 027 font foi.")
doc.add_paragraph("Des matrices de traçabilité, validateurs de schéma, modèles de transition, tests adversariaux et dossiers reproductibles peuvent servir d’aides. Aucun format ou outil n’est imposé ; ces éléments ne constituent pas seuls une preuve suffisante.")
doc.add_paragraph("Aucune identité ni relation globale entre REG-ARCH, REG, REG+ et TENIR-Gov n’est présumée. Toute revendication de réalisation d’une garantie est évaluée selon REG-ARCH-026, pour son périmètre et ses preuves exactes.")

# Candidate identity in package metadata.
doc.core_properties.title = "REG-ARCH v0.4.1 candidate — proposition amendée non ratifiée"
doc.core_properties.subject = "Amendements normatifs P0; REG-ARCH indépendant des implémentations"
doc.core_properties.comments = "Candidate locale produite pour audit contradictoire indépendant. Non ratifiée."
doc.save(OUT)
print(OUT)
