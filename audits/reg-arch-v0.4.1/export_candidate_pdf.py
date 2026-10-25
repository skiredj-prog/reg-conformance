"""Render the candidate DOCX structure to a readable PDF using ReportLab."""
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph as RP, Spacer, Table as RT, TableStyle, PageBreak

ROOT = Path(__file__).parent
src = ROOT / "REG-ARCH_v0.4.1-candidate.docx"
dst = ROOT / "REG-ARCH_v0.4.1-candidate.pdf"
doc = Document(src)
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="BodyFR", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.6, leading=11.2, spaceAfter=4))
styles.add(ParagraphStyle(name="Heading1FR", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=colors.HexColor("#17365D"), spaceBefore=12, spaceAfter=7, keepWithNext=True))
styles.add(ParagraphStyle(name="Heading2FR", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#1F4E79"), spaceBefore=8, spaceAfter=4, keepWithNext=True))
styles.add(ParagraphStyle(name="SmallCell", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.5, leading=8))
story = []

def safe(s):
    greek = {"τ": "tau", "Τ": "Tau", "Ω": "Omega", "ω": "omega", "λ": "lambda", "Λ": "Lambda", "γ": "gamma", "Γ": "Gamma", "δ": "delta", "Δ": "Delta", "μ": "mu", "Μ": "Mu", "σ": "sigma", "Σ": "Sigma", "α": "alpha", "Α": "Alpha", "β": "beta", "Β": "Beta", "ε": "epsilon", "Ε": "Epsilon"}
    for src, dst in greek.items():
        s = s.replace(src, dst)
    s = escape(s).replace("\n", "<br/>")
    return s.replace("→", "&#8594;").replace("—", "&#8212;").replace("–", "&#8211;").replace("É", "&#201;").replace("é", "&#233;").replace("è", "&#232;").replace("ê", "&#234;").replace("à", "&#224;").replace("ç", "&#231;").replace("ô", "&#244;").replace("û", "&#251;").replace("’", "&#8217;").replace("œ", "&#339;").replace("Œ", "&#338;")

def add_paragraph(p):
    text = p.text.strip()
    if not text:
        story.append(Spacer(1, 3))
        return
    style = p.style.name if p.style else "Normal"
    dest_style = "Heading1FR" if style.startswith("Heading 1") else "Heading2FR" if style.startswith("Heading 2") else "BodyFR"
    story.append(RP(safe(text), styles[dest_style]))

def add_table(t):
    data = [[RP(safe(c.text), styles["SmallCell"]) for c in row.cells] for row in t.rows]
    if not data:
        return
    usable = A4[0] - 34*mm
    n = max(1, len(data[0]))
    tbl = RT(data, colWidths=[usable/n]*n, repeatRows=1, hAlign="LEFT")
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#D9EAF7")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.HexColor("#17365D")),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#9EADBA")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 3), ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.extend([Spacer(1, 3), tbl, Spacer(1, 6)])

body = doc._element.body
for child in body.iterchildren():
    if isinstance(child, CT_P):
        add_paragraph(Paragraph(child, doc))
    elif isinstance(child, CT_Tbl):
        add_table(Table(child, doc))

def page(canvas, d):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#667788"))
    canvas.drawString(17*mm, 10*mm, "REG-ARCH v0.4.1 candidate — proposition non ratifiée")
    canvas.drawRightString(A4[0]-17*mm, 10*mm, f"{d.page}")
    canvas.restoreState()

SimpleDocTemplate(str(dst), pagesize=A4, rightMargin=17*mm, leftMargin=17*mm, topMargin=15*mm, bottomMargin=17*mm, title="REG-ARCH v0.4.1 candidate", author="REG-ARCH audit").build(story, onFirstPage=page, onLaterPages=page)
print(dst)
