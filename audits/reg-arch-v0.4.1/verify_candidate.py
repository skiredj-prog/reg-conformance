"""Consistency checks for the candidate package; not implementation conformance tests."""
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).parent
EXPECTED = {
    "source/REG-ARCH_v0_4.docx": "46752A0E1CD7AAE430742878C14C6B5C28FE1DA8167C9A81D78B23D9BAA7D3C2",
    "source/REG-ARCH_v0_4.pdf": "4BE9D68CA8DCE3557655905467C0975758E82AB2888F974B0B9439D59620B939",
}
for relative, digest in EXPECTED.items():
    actual = hashlib.sha256((ROOT / relative).read_bytes()).hexdigest().upper()
    assert actual == digest, f"source hash mismatch: {relative}"

schema = json.loads((ROOT / "trace-schema.json").read_text(encoding="utf-8"))
scenario_text = (ROOT / "adversarial-scenarios.yaml").read_text(encoding="utf-8")
match = re.search(r"required_trace_fields:\s*\[([^\]]+)\]", scenario_text)
assert match, "A-11 trace field list missing"
scenario_fields = [field.strip() for field in match.group(1).split(",")]
assert scenario_fields == schema["required_fields"], "YAML fields differ from normative trace schema"
scenario_ids = re.findall(r"^  - id: (A-\d+)$", scenario_text, re.M)
assert scenario_ids == [f"A-{n}" for n in range(11, 27)], f"unexpected scenario IDs: {scenario_ids}"

docx = Document(ROOT / "REG-ARCH_v0.4.1-candidate.docx")
doc_text = "\n".join([p.text for p in docx.paragraphs] + [c.text for t in docx.tables for r in t.rows for c in r.cells])
for clause in ("REG-ARCH-025", "REG-ARCH-026", "REG-ARCH-027"):
    assert clause in doc_text, f"missing normative clause {clause}"
for decision in (f"D-{n:02}" for n in range(1, 10)):
    assert decision in doc_text, f"missing open decision {decision}"
assert "Jusqu'à sa définition" not in doc_text
assert "Jusqu’à sa définition" not in doc_text
assert "transition informelle liée causalement" not in doc_text
assert "17. Amendements normatifs" in doc_text and "18. Décisions toujours ouvertes" in doc_text
assert "REG-ARCH-025 (§17)" in doc_text
assert "EXACTE de REG-ARCH" not in doc_text
assert "Appendix A — Informative, nonnormative" in doc_text
for forbidden in ("Incarnation exécutable de REG-ARCH", "REG Engine Code", "target_layer: \"REG Engine / TENIR-Gov\"", "traduction EXACTE"):
    assert forbidden not in doc_text, f"product-specific appendix assertion remains: {forbidden}"

pdf = PdfReader(ROOT / "REG-ARCH_v0.4.1-candidate.pdf")
pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
assert 8 <= len(pdf.pages) <= 15, f"unexpected PDF page count: {len(pdf.pages)}"
for required in ("REG-ARCH v0.4.1 candidate", "REG-ARCH-025", "REG-ARCH-026", "REG-ARCH-027", "Appendix A"):
    assert required in pdf_text, f"PDF missing {required}"
print(f"candidate package checks passed: 2 source hashes, trace schema/YAML parity, {len(scenario_ids)} scenarios, 3 clauses, 9 open decisions, {len(pdf.pages)}-page PDF")
